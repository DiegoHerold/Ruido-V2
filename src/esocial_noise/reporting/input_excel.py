from __future__ import annotations

from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
import shutil
import unicodedata

import pandas as pd
from openpyxl import load_workbook

from ..contracts import EmployeeInput
from ..utils import normalize_cpf


def _digits(value: object) -> str:
    return "".join(char for char in str(value or "") if char.isdigit())


def _header(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    return "".join(char for char in text if not unicodedata.combining(char)).strip().casefold()


def _is_completed(value: object) -> bool:
    return _header(value) == "concluido"


def _representation_type(value: object, row: int) -> str:
    normalized = " ".join(_header(value).split())
    if normalized in {"pessoa juridica", "juridica", "pj"}:
        return "pessoa_juridica"
    if normalized in {"pessoa fisica", "fisica", "pf"}:
        return "pessoa_fisica"
    raise ValueError(f"Tipo de procuração inválido na linha Excel {row}: informe 'Pessoa Jurídica' ou 'Pessoa Física'.")


def _expected_date(value: object, row: int) -> str:
    raw = "" if pd.isna(value) else str(value).strip()
    if not raw:
        return ""
    for pattern in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(raw, pattern).strftime("%d/%m/%Y")
        except ValueError:
            continue
    raise ValueError(f"Data da coluna Ano inválida na linha Excel {row}: informe dd/mm/aaaa.")


def _expected_intensity(value: object, row: int) -> str:
    raw = "" if pd.isna(value) else str(value).strip().replace(" ", "").replace(",", ".")
    if not raw:
        return ""
    try:
        number = Decimal(raw)
    except InvalidOperation as error:
        raise ValueError(f"Intensidade inválida na linha Excel {row}.") from error
    if number < 0:
        raise ValueError(f"Intensidade inválida na linha Excel {row}: use valor não negativo.")
    formatted = format(number.normalize(), "f")
    return formatted.rstrip("0").rstrip(".") if "." in formatted else formatted


def load_employees(path: Path) -> list[EmployeeInput]:
    if not path.exists() or path.suffix.lower() not in {".xlsx", ".xls", ".csv"}:
        raise ValueError("A entrada deve ser uma planilha existente (.xlsx, .xls ou .csv).")
    data = pd.read_csv(path, dtype=str) if path.suffix.lower() == ".csv" else pd.read_excel(path, dtype=str)
    data.columns = [_header(column) for column in data.columns]
    aliases = {
        "cpf": "cpf", "colaborador": "nome", "nome": "nome",
        "cnpj": "cnpj_representado", "cnpj representado": "cnpj_representado",
        "cpf representado": "cpf_representado", "cpf do representado": "cpf_representado",
        "tipo pessoa": "tipo_procuracao", "tipo de pessoa": "tipo_procuracao",
        "tipo procuracao": "tipo_procuracao", "tipo de procuracao": "tipo_procuracao",
        "pessoa": "tipo_procuracao", "perfil": "tipo_procuracao", "perfil procuracao": "tipo_procuracao",
        "ano": "data_planilha", "data": "data_planilha", "data inicio": "data_planilha",
        "intensidade": "intensidade_planilha",
        "status": "status",
    }
    data = data.rename(columns={source: destination for source, destination in aliases.items() if source in data.columns})
    required = {"cpf", "tipo_procuracao"}
    missing = required - set(data.columns)
    if missing:
        raise ValueError(f"Planilha inválida: colunas obrigatórias ausentes: {', '.join(sorted(missing))}.")
    employees: list[EmployeeInput] = []
    for index, row in data.iterrows():
        if _is_completed(row.get("status", "")):
            continue
        cpf = normalize_cpf(row["cpf"])
        if len(cpf) != 11:
            raise ValueError(f"CPF vazio ou inválido na linha Excel {index + 2}.")
        source_row = index + 2
        representation_type = _representation_type(row["tipo_procuracao"], source_row)
        document_column = "cnpj_representado" if representation_type == "pessoa_juridica" else "cpf_representado"
        represented_document = _digits(row.get(document_column, ""))
        expected_length = 14 if representation_type == "pessoa_juridica" else 11
        if len(represented_document) != expected_length:
            label = "CNPJ" if representation_type == "pessoa_juridica" else "CPF representado"
            raise ValueError(f"{label} vazio ou inválido na linha Excel {source_row}.")
        expected_date = _expected_date(row.get("data_planilha", ""), source_row)
        expected_intensity = _expected_intensity(row.get("intensidade_planilha", ""), source_row)
        employees.append(EmployeeInput(cpf, str(row.get("nome", "") or ""), source_row, representation_type, represented_document, expected_date, expected_intensity))
    return employees


def prepare_output_workbook(source_path: Path, output_dir: Path) -> Path:
    """Cria a planilha de saida como copia da entrada, preservando layout e filtros."""
    output_dir.mkdir(parents=True, exist_ok=True)
    destination = output_dir / f"{source_path.stem}_resultado{source_path.suffix}"
    shutil.copy2(source_path, destination)
    return destination


_DATE_HEADER_CANDIDATES = ("ano", "data", "data inicio")


def _fallback_unlocked_path(path: Path) -> Path:
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    return path.with_name(f"{path.stem}_recuperado_{stamp}{path.suffix}")


def mark_source_row_completed(path: Path, source_row: int, no_noise: bool, esocial_date: str = "", esocial_intensity: str = "", agent_code: str = "", log_text: str = "", status: str = "Concluido") -> Path:
    """Atualiza LOG/Status e, quando lidos do eSocial, Ano/Intensidade da linha concluída, preservando o restante da planilha."""
    if path.suffix.lower() == ".csv":
        data = pd.read_csv(path, dtype=str)
        if source_row - 2 >= len(data):
            raise ValueError(f"Linha Excel {source_row} inexistente para atualização de Status.")
        headers = {_header(column): column for column in data.columns}
        if "Status" not in data.columns:
            data["Status"] = ""
        if "LOG" not in data.columns:
            data["LOG"] = ""
        agent_column = headers.get("agente nocivo", "Agente Nocivo")
        if agent_column not in data.columns:
            data[agent_column] = ""
        date_column = next((headers[key] for key in _DATE_HEADER_CANDIDATES if key in headers), None) or "Ano"
        if date_column not in data.columns:
            data[date_column] = ""
        intensity_column = headers.get("intensidade", "Intensidade")
        if intensity_column not in data.columns:
            data[intensity_column] = ""
        data.loc[source_row - 2, "Status"] = status
        data.loc[source_row - 2, "LOG"] = log_text or ("Não há Agente de Ruído" if no_noise else "Agente de Ruído encontrado")
        if no_noise:
            data.loc[source_row - 2, "LOG"] = "Não há Agente de Ruído"
        if not no_noise and agent_code:
            data.loc[source_row - 2, agent_column] = agent_code
        if esocial_date:
            data.loc[source_row - 2, date_column] = esocial_date
        if esocial_intensity:
            data.loc[source_row - 2, intensity_column] = esocial_intensity
        data.loc[source_row - 2, "LOG"] = log_text or ("Não há Agente de Ruído" if no_noise else "Agente de Ruído encontrado")
        try:
            data.to_csv(path, index=False, encoding="utf-8-sig")
            return path
        except PermissionError:
            fallback = _fallback_unlocked_path(path)
            data.to_csv(fallback, index=False, encoding="utf-8-sig")
            return fallback
    if path.suffix.lower() != ".xlsx":
        raise ValueError("Atualização de Status é suportada apenas para .xlsx ou .csv.")
    workbook = load_workbook(path)
    worksheet = workbook.worksheets[0]
    columns = {_header(cell.value): cell.column for cell in worksheet[1] if cell.value is not None}
    status_column = columns.get("status")
    if not status_column:
        status_column = worksheet.max_column + 1
        worksheet.cell(row=1, column=status_column, value="Status")
    log_column = columns.get("log")
    if not log_column:
        log_column = worksheet.max_column + 1
        worksheet.cell(row=1, column=log_column, value="LOG")
    agent_column = columns.get("agente nocivo")
    if not agent_column:
        agent_column = worksheet.max_column + 1
        worksheet.cell(row=1, column=agent_column, value="Agente Nocivo")
    worksheet.cell(row=source_row, column=status_column, value=status)
    worksheet.cell(row=source_row, column=log_column, value=log_text or ("Não há Agente de Ruído" if no_noise else "Agente de Ruído encontrado"))
    if no_noise:
        worksheet.cell(row=source_row, column=log_column, value="Não há Agente de Ruído")
    if not no_noise and agent_code:
        worksheet.cell(row=source_row, column=agent_column, value=agent_code)
    if esocial_date:
        date_column = next((columns[key] for key in _DATE_HEADER_CANDIDATES if key in columns), None)
        if not date_column:
            date_column = worksheet.max_column + 1
            worksheet.cell(row=1, column=date_column, value="Ano")
        worksheet.cell(row=source_row, column=date_column, value=esocial_date)
    if esocial_intensity:
        intensity_column = columns.get("intensidade")
        if not intensity_column:
            intensity_column = worksheet.max_column + 1
            worksheet.cell(row=1, column=intensity_column, value="Intensidade")
        worksheet.cell(row=source_row, column=intensity_column, value=esocial_intensity)
    worksheet.cell(row=source_row, column=log_column, value=log_text or ("Não há Agente de Ruído" if no_noise else "Agente de Ruído encontrado"))
    try:
        workbook.save(path)
        return path
    except PermissionError:
        fallback = _fallback_unlocked_path(path)
        workbook.save(fallback)
        return fallback
